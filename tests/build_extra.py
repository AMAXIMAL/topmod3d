#!/usr/bin/env python3
"""Build reports for the remaining tools: primitive gallery + extended operator x modifier matrix."""
import json, os, html
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

BASE="/Users/ions/topmod-test"; SHOTS=os.path.join(BASE,"shots")
prims=json.load(open(os.path.join(BASE,"primitives.json")))
ops=json.load(open(os.path.join(BASE,"operators_extra.json")))

PLABEL={'cube':'Cube','tetrahedron':'Tetrahedron','octahedron':'Octahedron','icosahedron':'Icosahedron',
 'dodecahedron':'Dodecahedron','cuboctahedron':'Cuboctahedron','buckyball':'Buckyball','torus':'Torus',
 'cylinder':'Cylinder','mobius':'Möbius','spiral':'Spiral','polygon':'Polygon/Prism','rossler':'Rössler attractor',
 'sineWave':'Sine-wave surface','cymatics':'Cymatics plate','pseudospherical':'Pseudospherical','weave':'Woven tubes',
 'chainmail':'Chain mail','voronoi':'Voronoi','tree':'Tree','gem':'GEM matrix','maze':'Maze','maze3d':'Maze 3D'}
OLABEL={'move':'Move/Rotate/Scale','scale':'Uniform Scale','array':'Array','mirror':'Mirror','deleteSel':'Delete Selected',
 'flipNormals':'Flip Normals','heal':'Heal','quadRemesh':'Quad Remesh','simplify':'Simplify','planarize':'Planarize',
 'orient':'Orient to ground','cut':'Cut (plane)','pinch':'Pinch Extrude','bridge':'Bridge (2 faces)','pipe':'Pipes',
 'offset':'Offset Surface','sliceEdges':'Seam','slab':'Paper/Unfold (Slab)','gyroid':'Gyroid','sierpinski':'Sierpinski',
 'hybrid':'Schoen Hybrid','lattice':'Lattice (Cage)'}
MODES=['vert','edge','face']; ML={'vert':'Verts (1)','edge':'Edges (2)','face':'Faces (3)'}

VC={"PASS":"C6EFCE","PASS-vis":"CCECF5","NOEFFECT":"FFEB9C","DESTROYS":"FFD9B3","FAIL":"FFC7CE","MISSING":"D9D9D9","EMPTY":"FFC7CE"}
VF={"PASS":"1E7A34","PASS-vis":"125E7A","NOEFFECT":"8A6D00","DESTROYS":"9C5700","FAIL":"9C0006","MISSING":"555","EMPTY":"9C0006"}

# index ops
byop={}; oorder=[]
for r in ops:
    if r['op'] not in byop: byop[r['op']]={}; oorder.append(r['op'])
    byop[r['op']][r['mode']]=r

wb=Workbook()
# Sheet 1: primitives
ws=wb.active; ws.title="Primitives"
thin=Side(style="thin",color="CCCCCC"); bd=Border(left=thin,right=thin,top=thin,bottom=thin)
ws.append(["Primitive","id","Result","Topology (V / E / F / genus)","Screenshot"])
for c in ws[1]: c.font=Font(bold=True,color="FFFFFF"); c.fill=PatternFill("solid",fgColor="1F2A36")
for r in prims:
    row=[PLABEL.get(r['p'],r['p']), r['p'], r['verdict'], r.get('note','')]
    ws.append(row)
    rr=ws.max_row
    ws.cell(row=rr,column=3).fill=PatternFill("solid",fgColor=VC.get(r['verdict'],"FFF"))
    sc=ws.cell(row=rr,column=5,value="screenshot")
    sh=r.get('shot')
    if sh and os.path.exists(os.path.join(SHOTS,sh)):
        sc.hyperlink="shots/"+sh; sc.font=Font(color="0563C1",underline="single")
ws.column_dimensions["A"].width=20; ws.column_dimensions["B"].width=14
ws.column_dimensions["D"].width=34; ws.column_dimensions["E"].width=14

# Sheet 2: operators x modifier
ws2=wb.create_sheet("Operators x Modifier")
ws2.append(["Tool (operation)","op id",ML['vert'],ML['edge'],ML['face']])
for c in ws2[1]: c.font=Font(bold=True,color="FFFFFF"); c.fill=PatternFill("solid",fgColor="1F2A36"); c.alignment=Alignment(horizontal="center",wrap_text=True)
for op in oorder:
    ws2.append([OLABEL.get(op,op),op,"","",""])
    rr=ws2.max_row
    for mi,m in enumerate(MODES):
        rec=byop[op].get(m); v=rec['verdict'] if rec else 'MISSING'
        cell=ws2.cell(row=rr,column=3+mi,value=v+"\n"+(rec.get('note','') if rec else ''))
        cell.fill=PatternFill("solid",fgColor=VC.get(v,"FFF")); cell.font=Font(color=VF.get(v,"000"),size=9)
        cell.alignment=Alignment(horizontal="center",vertical="center",wrap_text=True)
        if rec and rec.get('shot') and os.path.exists(os.path.join(SHOTS,rec['shot'])): cell.hyperlink="shots/"+rec['shot']
    ws2.row_dimensions[rr].height=34
ws2.column_dimensions["A"].width=20; ws2.column_dimensions["B"].width=12
for c in ["C","D","E"]: ws2.column_dimensions[c].width=24
ws2.freeze_panes="A2"

# Legend
ws3=wb.create_sheet("Legend")
leg=[("PASS","Tool ran and changed mesh topology (V/E/F)."),
     ("PASS-vis","Topology unchanged but the rendered geometry changed (move/scale/flip/offset/etc.) — detected by image diff."),
     ("NOEFFECT","Ran with no detectable change — usually needs a manual pick/gesture (Bridge=2 faces, Lattice=host shape, Orient=face pick) or a different selection mode."),
     ("DESTROYS","Emptied the mesh (Delete Selected with all selected)."),
     ("FAIL","Threw a JavaScript error. NOTE: 'Paper/Unfold (Slab)' fails in all 3 modes — a real bug to investigate."),
     ("MISSING","Button not found.")]
ws3.append(["Verdict","Meaning"])
for c in ws3[1]: c.font=Font(bold=True)
for k,d in leg:
    ws3.append([k,d]); ws3.cell(row=ws3.max_row,column=1).fill=PatternFill("solid",fgColor=VC[k])
ws3.column_dimensions["A"].width=12; ws3.column_dimensions["B"].width=100

xlsx=os.path.join(BASE,"TOPMOD3D_all_other_tools.xlsx"); wb.save(xlsx)

# HTML
def img(sh,t):
    return f'<a href="shots/{sh}" target="_blank"><img src="shots/{sh}" loading="lazy" title="{t}"></a>' if sh and os.path.exists(os.path.join(SHOTS,sh)) else '<div class=noimg>no image</div>'
pcards="".join(f'<div class=card>{img(r.get("shot"),r["p"])}<div class=cl>{PLABEL.get(r["p"],r["p"])}</div><div class=note>{html.escape(r.get("note",""))}</div></div>' for r in prims)
def opcell(op,m):
    rec=byop[op].get(m)
    if not rec: return '<td><div class="v MISSING">—</div></td>'
    v=rec['verdict']
    return f'<td>{img(rec.get("shot"),op+" / "+m)}<div class="v {v}">{v}</div><div class=note>{html.escape(rec.get("note",""))}</div></td>'
orows="".join('<tr><th class=op><div class=ol>'+html.escape(OLABEL.get(op,op))+'</div><code>'+op+'</code></th>'+"".join(opcell(op,m) for m in MODES)+'</tr>' for op in oorder)
from collections import Counter
pc=Counter(r['verdict'] for r in ops)
summ=" &middot; ".join(f'{k} {v}' for k,v in sorted(pc.items()))
doc=f"""<!doctype html><meta charset=utf-8><title>TOPMOD 3D — all other tools</title>
<style>
body{{background:#0d1420;color:#cfe8ff;font:14px system-ui,sans-serif;margin:0;padding:22px}}
h1{{font-size:20px;margin:0 0 2px}} h2{{font-size:16px;margin:26px 0 10px;color:#9fd4ff}}
.sub{{color:#7fa6c4;font-size:12px;margin:0 0 14px}}
.grid{{display:grid;grid-template-columns:repeat(auto-fill,minmax(190px,1fr));gap:12px}}
.card{{background:#121a24;border:1px solid #21303f;border-radius:8px;padding:8px;text-align:center}}
.card img{{width:100%;border-radius:6px;background:#000}}
.cl{{font-weight:700;margin-top:6px;color:#fff;font-size:13px}}
table{{border-collapse:collapse;width:100%}} th,td{{border:1px solid #21303f;padding:8px;vertical-align:top;text-align:center}}
th.op{{text-align:left;width:160px;background:#121a24}} th.op .ol{{font-weight:700;color:#fff}} th.op code{{color:#6f93ad;font-size:11px}}
thead th{{background:#16202c}} td{{width:30%}} td img{{width:100%;max-width:280px;border-radius:6px;background:#000}}
.v{{font-weight:700;font-size:11px;padding:2px 6px;border-radius:4px;display:inline-block;margin:6px 0 2px}}
.note{{color:#88a0b4;font-size:11px;font-family:ui-monospace,monospace}}
.PASS{{background:#1e7a34;color:#fff}} .PASS-vis{{background:#125e7a;color:#fff}} .NOEFFECT{{background:#8a6d00;color:#fff}}
.DESTROYS{{background:#9c5700;color:#fff}} .FAIL{{background:#9c0006;color:#fff}} .MISSING{{background:#555;color:#fff}}
.noimg{{color:#566;padding:40px 0;font-size:11px}}
</style>
<h1>TOPMOD 3D — all other tools</h1>
<p class=sub>Coverage of the remaining tools not in the earlier matrices: 23 shape primitives + 22 operators × 3 modifiers.</p>
<h2>Shape primitives (23) — all generate successfully</h2>
<div class=grid>{pcards}</div>
<h2>Remaining operators × modifier (66 cells) — {summ}</h2>
<p class=sub>Cube → modifier → Select All → tool. PASS = topology changed · PASS-vis = geometry moved (image diff) · NOEFFECT = needs manual pick/gesture · FAIL = error.</p>
<table><thead><tr><th class=op>Tool \\ Modifier</th><th>{ML['vert']}</th><th>{ML['edge']}</th><th>{ML['face']}</th></tr></thead><tbody>{orows}</tbody></table>
"""
open(os.path.join(BASE,"report_all_other.html"),"w").write(doc)
print("XLSX:",xlsx); print("HTML:",os.path.join(BASE,"report_all_other.html"))
print("prims:",len(prims),"ops:",len(ops),"opcounts:",dict(pc))
