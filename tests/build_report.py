#!/usr/bin/env python3
"""Build TOPMOD 3D tool x modifier compatibility report from sweep results."""
import json, os, html
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

BASE = "/Users/ions/topmod-test"
SHOTS = os.path.join(BASE, "shots")
res = json.load(open(os.path.join(BASE, "results.json")))

MODES = ["vert", "edge", "face"]
MODE_LABEL = {"vert": "Verts (1)", "edge": "Edges (2)", "face": "Faces (3)"}

# index: op -> mode -> record
ops_order = []
by_op = {}
op_label = {}
for r in res:
    op = r["op"]
    if op not in by_op:
        by_op[op] = {}
        ops_order.append(op)
    by_op[op][r["mode"]] = r
    op_label[op] = r.get("opLabel", op)

VCOLOR = {
    "PASS":     "C6EFCE",  # green
    "NOEFFECT": "FFEB9C",  # amber
    "DESTROYS": "FFD9B3",  # orange
    "FAIL":     "FFC7CE",  # red
    "MISSING":  "D9D9D9",  # grey
    "UNVERIFIED": "DCE6F1",  # light blue
}
VFONT = {
    "PASS":     "1E7A34", "NOEFFECT": "8A6D00", "DESTROYS": "9C5700",
    "FAIL":     "9C0006", "MISSING":  "555555", "UNVERIFIED": "2E5B8A",
}

# ---------- XLSX ----------
wb = Workbook()
ws = wb.active
ws.title = "Tool x Modifier"
thin = Side(style="thin", color="CCCCCC")
border = Border(left=thin, right=thin, top=thin, bottom=thin)
hdr_fill = PatternFill("solid", fgColor="1F2A36")
hdr_font = Font(bold=True, color="FFFFFF", size=11)
center = Alignment(horizontal="center", vertical="center", wrap_text=True)
left = Alignment(horizontal="left", vertical="center", wrap_text=True)

# Title
ws.merge_cells("A1:H1")
ws["A1"] = "TOPMOD 3D — Tool x Modifier compatibility (start mesh: Cube, Select All applied)"
ws["A1"].font = Font(bold=True, size=13)
ws.row_dimensions[1].height = 22

# header row 2
headers = ["Tool (operation)", "op id"]
for m in MODES:
    headers.append(MODE_LABEL[m])
headers += ["Screens (vert / edge / face)"]
for ci, h in enumerate(headers, 1):
    c = ws.cell(row=2, column=ci, value=h)
    c.fill = hdr_fill; c.font = hdr_font; c.alignment = center; c.border = border

row = 3
for op in ops_order:
    ws.cell(row=row, column=1, value=op_label[op]).alignment = left
    ws.cell(row=row, column=2, value=op).alignment = left
    for mi, m in enumerate(MODES):
        rec = by_op[op].get(m)
        col = 3 + mi
        v = rec["verdict"] if rec else "MISSING"
        note = rec.get("note", "") if rec else ""
        cell = ws.cell(row=row, column=col, value=f"{v}\n{note}")
        cell.fill = PatternFill("solid", fgColor=VCOLOR.get(v, "FFFFFF"))
        cell.font = Font(color=VFONT.get(v, "000000"), size=9)
        cell.alignment = center; cell.border = border
    # screenshot hyperlinks
    links = []
    for m in MODES:
        rec = by_op[op].get(m)
        if rec and rec.get("shot"):
            p = os.path.join(SHOTS, rec["shot"])
            if os.path.exists(p):
                links.append(rec["shot"])
    lc = ws.cell(row=row, column=6, value="open: " + " | ".join(MODES))
    # link to the face shot if present else first
    target = None
    for m in ["face", "edge", "vert"]:
        rec = by_op[op].get(m)
        if rec and rec.get("shot") and os.path.exists(os.path.join(SHOTS, rec["shot"])):
            target = "shots/" + rec["shot"]; break
    if target:
        lc.hyperlink = target; lc.font = Font(color="0563C1", underline="single", size=9)
    lc.alignment = left
    for ci in range(1, 7):
        ws.cell(row=row, column=ci).border = border
    ws.row_dimensions[row].height = 34
    row += 1

ws.column_dimensions["A"].width = 24
ws.column_dimensions["B"].width = 14
for c in ["C", "D", "E"]:
    ws.column_dimensions[c].width = 22
ws.column_dimensions["F"].width = 26
ws.freeze_panes = "A3"

# Legend sheet
ws2 = wb.create_sheet("Legend")
legend = [
    ("PASS", "Tool ran and changed the mesh into a valid result in this selection mode."),
    ("NOEFFECT", "Tool ran without error but produced no change (not applicable in this mode, or needs a different selection)."),
    ("DESTROYS", "Tool ran but emptied/destroyed the mesh (e.g. Delete with everything selected)."),
    ("FAIL", "Tool threw a JavaScript error in this mode."),
    ("UNVERIFIED", "Selection-only tool in Verts/Edges mode: the harness measures face-selection only, so the effect could not be captured automatically. Works interactively."),
    ("MISSING", "Tool button not found in this run."),
]
ws2.cell(row=1, column=1, value="Verdict").font = Font(bold=True)
ws2.cell(row=1, column=2, value="Meaning").font = Font(bold=True)
for i, (k, d) in enumerate(legend, 2):
    a = ws2.cell(row=i, column=1, value=k)
    a.fill = PatternFill("solid", fgColor=VCOLOR[k]); a.font = Font(color=VFONT[k], bold=True)
    ws2.cell(row=i, column=2, value=d)
ws2.column_dimensions["A"].width = 14
ws2.column_dimensions["B"].width = 90

# Raw sheet
ws3 = wb.create_sheet("Raw data")
cols = ["idx", "mode", "op", "opLabel", "verdict", "note", "param", "shot"]
for ci, h in enumerate(cols, 1):
    ws3.cell(row=1, column=ci, value=h).font = Font(bold=True)
for ri, r in enumerate(res, 2):
    for ci, k in enumerate(cols, 1):
        ws3.cell(row=ri, column=ci, value=str(r.get(k, "")))

xlsx_path = os.path.join(BASE, "TOPMOD3D_tool_modifier_matrix.xlsx")
wb.save(xlsx_path)

# ---------- HTML gallery ----------
def cellhtml(rec):
    if not rec:
        return '<td class="m"><div class="v MISSING">MISSING</div></td>'
    v = rec["verdict"]; shot = rec.get("shot", "")
    img = f'<a href="shots/{shot}" target="_blank"><img src="shots/{shot}" loading="lazy"></a>' if shot and os.path.exists(os.path.join(SHOTS, shot)) else '<div class="noimg">no image</div>'
    return f'<td class="m"><div class="v {v}">{v}</div>{img}<div class="note">{html.escape(rec.get("note",""))}</div></td>'

rows_html = []
for op in ops_order:
    cells = "".join(cellhtml(by_op[op].get(m)) for m in MODES)
    rows_html.append(f'<tr><th class="op"><div class="ol">{html.escape(op_label[op])}</div><code>{op}</code></th>{cells}</tr>')

counts = {}
for r in res:
    counts[r["verdict"]] = counts.get(r["verdict"], 0) + 1
summary = " &middot; ".join(f'<b class="{k}">{k}</b> {v}' for k, v in sorted(counts.items()))

doc = f"""<!doctype html><meta charset="utf-8"><title>TOPMOD 3D — Tool x Modifier matrix</title>
<style>
body{{background:#0d1420;color:#cfe8ff;font:14px -apple-system,system-ui,sans-serif;margin:0;padding:24px}}
h1{{font-size:20px;margin:0 0 4px}} .sub{{color:#7fa6c4;margin:0 0 16px;font-size:13px}}
.summary{{margin:0 0 18px;font-size:13px}}
table{{border-collapse:collapse;width:100%}}
th,td{{border:1px solid #21303f;padding:8px;vertical-align:top}}
thead th{{background:#16202c;position:sticky;top:0;z-index:2}}
th.op{{text-align:left;width:170px;background:#121a24}}
th.op .ol{{font-weight:700;color:#fff}} th.op code{{color:#6f93ad;font-size:11px}}
td.m{{width:31%;text-align:center}}
td.m img{{width:100%;max-width:300px;border-radius:6px;border:1px solid #20303f;background:#000}}
.v{{font-weight:700;font-size:11px;letter-spacing:.5px;padding:2px 6px;border-radius:4px;display:inline-block;margin-bottom:6px}}
.note{{color:#88a0b4;font-size:11px;margin-top:5px;font-family:ui-monospace,monospace}}
.PASS{{background:#1e7a34;color:#fff}} .NOEFFECT{{background:#8a6d00;color:#fff}}
.DESTROYS{{background:#9c5700;color:#fff}} .FAIL{{background:#9c0006;color:#fff}} .MISSING{{background:#555;color:#fff}}
.UNVERIFIED{{background:#2e5b8a;color:#fff}}
.noimg{{color:#566;font-size:11px;padding:40px 0}}
</style>
<h1>TOPMOD 3D — Tool &times; Modifier compatibility matrix</h1>
<p class="sub">Each cell = create a <b>Cube</b> &rarr; switch to the modifier (selection mode) &rarr; <b>Select All</b> &rarr; apply the tool. {len(res)} combinations.</p>
<p class="summary">{summary}</p>
<table><thead><tr><th class="op">Tool \\ Modifier</th><th>{MODE_LABEL['vert']}</th><th>{MODE_LABEL['edge']}</th><th>{MODE_LABEL['face']}</th></tr></thead>
<tbody>{''.join(rows_html)}</tbody></table>
"""
open(os.path.join(BASE, "report.html"), "w").write(doc)
print("XLSX:", xlsx_path)
print("HTML:", os.path.join(BASE, "report.html"))
print("counts:", counts)
