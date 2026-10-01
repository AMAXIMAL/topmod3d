#!/usr/bin/env python3
"""Primitive x Operator matrix (23 primitives x 16 operators, Faces mode)."""
import json, os, html
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

BASE="/Users/ions/topmod-test"; SHOTS=os.path.join(BASE,"shots")
res=json.load(open(os.path.join(BASE,"primitive_x_operator.json")))
PRIMS=['cube','tetrahedron','octahedron','icosahedron','dodecahedron','cuboctahedron','buckyball','torus','cylinder','mobius','spiral','polygon','rossler','sineWave','cymatics','pseudospherical','weave','chainmail','voronoi','tree','gem','maze','maze3d']
OPS=['triangulate','stellate','doubleStellate','extrude','dome','crust','wireframe','menger','catmull','doosabin','loop','honeycomb','dual','truncate','pentagonalize','root3']
PL={'cube':'Cube','tetrahedron':'Tetra','octahedron':'Octa','icosahedron':'Icosa','dodecahedron':'Dodeca','cuboctahedron':'Cubocta','buckyball':'Buckyball','torus':'Torus','cylinder':'Cylinder','mobius':'Möbius','spiral':'Spiral','polygon':'Polygon','rossler':'Rössler','sineWave':'SineWave','cymatics':'Cymatics','pseudospherical':'Pseudosph','weave':'Weave','chainmail':'Chainmail','voronoi':'Voronoi','tree':'Tree','gem':'GEM','maze':'Maze','maze3d':'Maze3D'}
OL={'triangulate':'Triangulate','stellate':'Stellate','doubleStellate':'Dbl Stellate','extrude':'Extrude','dome':'Dome','crust':'Crust','wireframe':'Wireframe','menger':'Menger','catmull':'Catmull','doosabin':'Doo-Sabin','loop':'Loop','honeycomb':'Honeycomb','dual':'Dual','truncate':'Truncate','pentagonalize':'Pentagonal','root3':'√3'}
grid={(r['prim'],r['op']):r for r in res}
VC={"PASS":"C6EFCE","NOEFFECT":"FFEB9C","SKIP-heavy":"E4D9F2","DESTROYS":"FFD9B3","FAIL":"FFC7CE","MISSING-prim":"D9D9D9","MISSING-op":"D9D9D9"}
VF={"PASS":"1E7A34","NOEFFECT":"8A6D00","SKIP-heavy":"5B3A8A","DESTROYS":"9C5700","FAIL":"9C0006"}
AB={"PASS":"✓","NOEFFECT":"·","SKIP-heavy":"hvy","DESTROYS":"✗","FAIL":"ERR"}

wb=Workbook(); ws=wb.active; ws.title="Primitive x Operator"
thin=Side(style="thin",color="CCCCCC"); bd=Border(left=thin,right=thin,top=thin,bottom=thin)
ws.cell(row=1,column=1,value="Primitive \\ Operator  (Cube-style: create primitive → Select All faces → apply op)").font=Font(bold=True,size=12)
for ci,op in enumerate(OPS,2):
    c=ws.cell(row=2,column=ci,value=OL[op]); c.alignment=Alignment(text_rotation=90,vertical="bottom",horizontal="center"); c.font=Font(bold=True,size=9); c.border=bd
ws.cell(row=2,column=1,value="prim \\ op").font=Font(bold=True)
ws.row_dimensions[2].height=78
for ri,p in enumerate(PRIMS,3):
    h=ws.cell(row=ri,column=1,value=PL[p]); h.font=Font(bold=True,size=9); h.border=bd
    for ci,op in enumerate(OPS,2):
        rec=grid.get((p,op)); v=rec['verdict'] if rec else 'MISSING-op'
        cell=ws.cell(row=ri,column=ci,value=AB.get(v,v))
        cell.fill=PatternFill("solid",fgColor=VC.get(v,"FFFFFF")); cell.font=Font(color=VF.get(v,"000000"),bold=True,size=10)
        cell.alignment=Alignment(horizontal="center",vertical="center"); cell.border=bd
        if rec and rec.get('shot') and os.path.exists(os.path.join(SHOTS,rec['shot'])): cell.hyperlink="shots/"+rec['shot']
ws.column_dimensions["A"].width=12
for ci in range(2,len(OPS)+2): ws.column_dimensions[get_column_letter(ci)].width=6
ws.freeze_panes="B3"

ws2=wb.create_sheet("Legend & notes")
for i,(k,d) in enumerate([
 ("PASS ✓","Operator applied to this primitive and changed its topology — valid result."),
 ("NOEFFECT ·","Ran with no change (already in that form, or op needs a selection/condition this primitive doesn't meet)."),
 ("SKIP-heavy hvy","Skipped automatically: projected faces (base × op multiplier) exceed the safe ~6000 budget. These combos FREEZE the app (synchronous, no progress) — see error log. 128 combos affected."),
 ("DESTROYS ✗","Emptied the mesh."),("FAIL ERR","JavaScript error."),
],1):
    ws2.cell(row=i,column=1,value=k).fill=PatternFill("solid",fgColor=VC.get(k.split()[0],"FFFFFF"))
    ws2.cell(row=i,column=2,value=d)
ws2.column_dimensions["A"].width=16; ws2.column_dimensions["B"].width=100

ws3=wb.create_sheet("Raw");
for ci,h in enumerate(["prim","op","verdict","note","err","shot"],1): ws3.cell(row=1,column=ci,value=h).font=Font(bold=True)
for ri,r in enumerate(res,2):
    for ci,k in enumerate(["prim","op","verdict","note","err","shot"],1): ws3.cell(row=ri,column=ci,value=str(r.get(k,"")))

xlsx=os.path.join(BASE,"TOPMOD3D_primitive_x_operator.xlsx"); wb.save(xlsx)

def img(sh,t): return f'<a href="shots/{sh}" target="_blank"><img src="shots/{sh}" loading="lazy" title="{t}"></a>' if sh and os.path.exists(os.path.join(SHOTS,sh)) else ''
def cell(p,op):
    rec=grid.get((p,op));
    if not rec: return '<td></td>'
    v=rec['verdict']
    return f'<td class="{v}">{img(rec.get("shot"),p+" / "+op)}<div class=vt>{AB.get(v,v)}</div></td>'
head="".join(f'<th><div class=rot>{OL[op]}</div></th>' for op in OPS)
body="".join('<tr><th class=rh>'+PL[p]+'</th>'+"".join(cell(p,op) for op in OPS)+'</tr>' for p in PRIMS)
from collections import Counter
C=Counter(r['verdict'] for r in res); summ=" · ".join(f'{k} {v}' for k,v in sorted(C.items()))
doc=f"""<!doctype html><meta charset=utf-8><title>TOPMOD 3D — primitive x operator</title>
<style>
body{{background:#0d1420;color:#cfe8ff;font:13px system-ui,sans-serif;margin:0;padding:20px}}
h1{{font-size:19px;margin:0 0 2px}} .sub{{color:#7fa6c4;font-size:12px;margin:0 0 12px}}
table{{border-collapse:collapse}} td,th{{border:1px solid #21303f}}
th.rh{{background:#121a24;text-align:right;padding:4px 8px;font-size:11px;color:#fff;white-space:nowrap}}
thead th{{background:#16202c;height:110px;vertical-align:bottom}} .rot{{writing-mode:vertical-rl;transform:rotate(180deg);font-size:11px;padding:6px 2px}}
td{{width:70px;height:70px;padding:0;text-align:center;position:relative}} td img{{width:68px;height:68px;object-fit:cover;display:block}}
.vt{{position:absolute;top:1px;left:2px;font-weight:700;font-size:10px;text-shadow:0 0 3px #000}}
td.PASS{{outline:2px solid #1e7a34;outline-offset:-2px}} td.SKIP-heavy{{background:#2a1f3a;color:#caa6ff}} td.NOEFFECT{{outline:2px solid #8a6d00;outline-offset:-2px}}
td.FAIL,td.DESTROYS{{outline:2px solid #9c0006;outline-offset:-2px}}
.legend{{margin:10px 0;font-size:12px;color:#9fb4c8}}
</style>
<h1>TOPMOD 3D — Primitive × Operator</h1>
<p class=sub>Each of 23 primitives → Select All faces → apply each of 16 operators. {len(res)} combinations. <b>Normals shader</b> (face colour = surface-normal direction). No crashes; <b>hvy</b> = skipped to avoid the known UI-freeze on high-poly primitives.</p>
<p class=legend>{summ}</p>
<table><thead><tr><th class=rh>prim \\ op</th>{head}</tr></thead><tbody>{body}</tbody></table>
"""
open(os.path.join(BASE,"report_pxo.html"),"w").write(doc)
print("XLSX:",xlsx); print("HTML:",os.path.join(BASE,"report_pxo.html")); print("counts:",dict(C))
