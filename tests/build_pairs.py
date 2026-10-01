#!/usr/bin/env python3
"""Build TOPMOD 3D pairwise tool-sequence matrix (A then B, Faces mode)."""
import json, os, html
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

BASE = "/Users/ions/topmod-test"
SHOTS = os.path.join(BASE, "shots")
res = json.load(open(os.path.join(BASE, "pairs.json")))

TOOLS = ['triangulate','stellate','doubleStellate','extrude','dome','crust','wireframe',
         'menger','catmull','doosabin','loop','honeycomb','dual','truncate','pentagonalize','root3']
LABEL = {'triangulate':'Triangulate','stellate':'Stellate','doubleStellate':'Double Stellate',
         'extrude':'Extrude','dome':'Dome','crust':'Crust','wireframe':'Wireframe','menger':'Menger',
         'catmull':'Catmull-Clark','doosabin':'Doo-Sabin','loop':'Loop','honeycomb':'Honeycomb',
         'dual':'Dual','truncate':'Truncate','pentagonalize':'Pentagonalize','root3':'√3'}

grid = {}
for r in res:
    grid[(r['A'], r['B'])] = r

VCOLOR = {"PASS":"C6EFCE","NOEFFECT":"FFEB9C","DESTROYS":"FFD9B3","FAIL":"FFC7CE",
          "SKIP-heavy":"E4DFEC","A-DESTROYS":"F2DCDB","MISSING":"D9D9D9"}
VFONT  = {"PASS":"1E7A34","NOEFFECT":"8A6D00","DESTROYS":"9C5700","FAIL":"9C0006",
          "SKIP-heavy":"5B4A8A","A-DESTROYS":"9C0006","MISSING":"555555"}
ABBR   = {"PASS":"✓","NOEFFECT":"·","DESTROYS":"✗","FAIL":"ERR",
          "SKIP-heavy":"hvy","A-DESTROYS":"A✗","MISSING":"?"}

wb = Workbook(); ws = wb.active; ws.title = "Pairwise (A then B)"
thin = Side(style="thin", color="CCCCCC"); border = Border(left=thin,right=thin,top=thin,bottom=thin)
center = Alignment(horizontal="center", vertical="center", wrap_text=True)

ws.merge_cells(start_row=1,start_column=1,end_row=1,end_column=len(TOOLS)+1)
ws["A1"] = "TOPMOD 3D — Pairwise tool sequence (Cube → A → Select All → B, Faces mode).  Rows = first tool A, Columns = second tool B."
ws["A1"].font = Font(bold=True, size=12)

ws.cell(row=2, column=1, value="A \\ B").font = Font(bold=True)
ws.cell(row=2, column=1).alignment = center
for ci, B in enumerate(TOOLS, 2):
    c = ws.cell(row=2, column=ci, value=LABEL[B])
    c.font = Font(bold=True, size=9); c.alignment = Alignment(horizontal="center", vertical="bottom", wrap_text=True, text_rotation=90)
    c.border = border
ws.row_dimensions[2].height = 90

for ri, A in enumerate(TOOLS, 3):
    h = ws.cell(row=ri, column=1, value=LABEL[A]); h.font = Font(bold=True, size=9)
    h.alignment = Alignment(horizontal="left", vertical="center"); h.border = border
    for ci, B in enumerate(TOOLS, 2):
        rec = grid.get((A, B))
        v = rec['verdict'] if rec else 'MISSING'
        cell = ws.cell(row=ri, column=ci, value=ABBR.get(v, v))
        cell.fill = PatternFill("solid", fgColor=VCOLOR.get(v, "FFFFFF"))
        cell.font = Font(color=VFONT.get(v, "000000"), bold=True, size=11)
        cell.alignment = center; cell.border = border
        if rec and rec.get('shot') and os.path.exists(os.path.join(SHOTS, rec['shot'])):
            cell.hyperlink = "shots/" + rec['shot']
        if rec:
            cell.comment = None
ws.column_dimensions["A"].width = 16
for ci in range(2, len(TOOLS)+2):
    from openpyxl.utils import get_column_letter
    ws.column_dimensions[get_column_letter(ci)].width = 5
ws.freeze_panes = "B3"

# legend
ws2 = wb.create_sheet("Legend")
leg = [("PASS ✓","B applied successfully to A's result (mesh changed, valid)."),
       ("NOEFFECT ·","B ran but produced no change on A's result."),
       ("DESTROYS ✗","B emptied the mesh."),
       ("SKIP-heavy hvy","A produced a heavy mesh (>2500 faces); B skipped to avoid freeze."),
       ("A-DESTROYS A✗","A itself emptied the mesh, so B could not run."),
       ("FAIL ERR","JavaScript error.")]
ws2.cell(row=1,column=1,value="Code").font=Font(bold=True); ws2.cell(row=1,column=2,value="Meaning").font=Font(bold=True)
for i,(k,d) in enumerate(leg,2):
    ws2.cell(row=i,column=1,value=k); ws2.cell(row=i,column=2,value=d)
ws2.column_dimensions["A"].width=18; ws2.column_dimensions["B"].width=80

# raw
ws3 = wb.create_sheet("Raw")
cols=["idx","A","B","verdict","note","shot"]
for ci,h in enumerate(cols,1): ws3.cell(row=1,column=ci,value=h).font=Font(bold=True)
for ri,r in enumerate(res,2):
    for ci,k in enumerate(cols,1): ws3.cell(row=ri,column=ci,value=str(r.get(k,"")))

xlsx=os.path.join(BASE,"TOPMOD3D_pairwise_sequence_matrix.xlsx"); wb.save(xlsx)

# HTML grid
def cellh(A,B):
    rec=grid.get((A,B));
    if not rec: return '<td><div class="v MISSING">?</div></td>'
    v=rec['verdict']; shot=rec.get('shot','')
    img=f'<a href="shots/{shot}" target="_blank"><img src="shots/{shot}" loading="lazy" title="{A} then {B}: {v}"></a>' if shot and os.path.exists(os.path.join(SHOTS,shot)) else ''
    return f'<td class="{v}">{img}<div class="vt">{ABBR.get(v,v)}</div></td>'
head="".join(f'<th><div class="rot">{LABEL[B]}</div></th>' for B in TOOLS)
body="".join('<tr><th class="rh">'+LABEL[A]+'</th>'+"".join(cellh(A,B) for B in TOOLS)+'</tr>' for A in TOOLS)
counts={}
for r in res: counts[r['verdict']]=counts.get(r['verdict'],0)+1
summ=" &middot; ".join(f'{k} {v}' for k,v in sorted(counts.items()))
doc=f"""<!doctype html><meta charset=utf-8><title>TOPMOD 3D pairwise sequences</title>
<style>
body{{background:#0d1420;color:#cfe8ff;font:13px system-ui,sans-serif;margin:0;padding:20px}}
h1{{font-size:19px;margin:0 0 4px}} .sub{{color:#7fa6c4;font-size:12px;margin:0 0 14px}}
table{{border-collapse:collapse}} td,th{{border:1px solid #21303f}}
th.rh{{background:#121a24;text-align:right;padding:4px 8px;font-size:11px;color:#fff;white-space:nowrap}}
thead th{{background:#16202c;height:120px;vertical-align:bottom}}
.rot{{writing-mode:vertical-rl;transform:rotate(180deg);font-size:11px;padding:6px 2px;color:#cfe8ff}}
td{{width:74px;height:74px;padding:0;position:relative;text-align:center;vertical-align:middle}}
td img{{width:72px;height:72px;object-fit:cover;display:block}}
td .vt{{position:absolute;top:1px;left:2px;font-weight:700;font-size:11px;text-shadow:0 0 3px #000}}
td.PASS{{outline:2px solid #1e7a34;outline-offset:-2px}} td.NOEFFECT{{outline:2px solid #8a6d00;outline-offset:-2px}}
td.DESTROYS,td.FAIL,td["A-DESTROYS"]{{outline:2px solid #9c0006;outline-offset:-2px}}
td.MISSING,td.v{{background:#1a2430;color:#566}}
.legend{{margin:14px 0;font-size:12px;color:#9fb4c8}}
.legend b{{color:#fff}}
</style>
<h1>TOPMOD 3D — Pairwise tool sequences</h1>
<p class=sub>Cube → tool <b>A</b> (row) → Select All → tool <b>B</b> (column), Faces mode. {len(res)} sequences. ✓ works · no-effect ✗ destroys.</p>
<p class=legend>{summ}</p>
<table><thead><tr><th class=rh>A \\ B</th>{head}</tr></thead><tbody>{body}</tbody></table>
"""
open(os.path.join(BASE,"report_pairs.html"),"w").write(doc)
print("XLSX:",xlsx); print("HTML:",os.path.join(BASE,"report_pairs.html")); print("counts:",counts)
